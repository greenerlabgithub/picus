import logging
import azure.functions as func

import sys
import json
import base64
import cv2
import numpy as np
import openpyxl
import os
from openpyxl.styles import Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

# --------------------------------
# (1) 런북 코드 재활용
# --------------------------------

color_ranges = {
    "검정":  ((0,   0,   0),   (10, 255,  50)),
    "갈색":  ((10,  80,  5),   (30, 255, 255)),
    "초록":  ((60,  50,  50),  (85, 255, 255)),
    "보라":  ((130, 50,  50),  (170, 255, 255)),
    "파랑":  ((90,  50,  50),  (114, 255, 255)),

    "나무둘레": ((115, 50, 50),  (129, 255, 255)),
    "나무표시": ((32,  50, 50),  (59, 255, 255)),
    "나무표시2":((3,  240,150),  (10, 255, 255))
}

grade_count = {"A":0, "B":0, "C":0, "D":0, "E":0}

def calc_grade(ratio: float) -> str:
    if 0 <= ratio < 1:
        return "A"
    elif 1 <= ratio <= 19:
        return "B"
    elif 20 <= ratio <= 39:
        return "C"
    elif 40 <= ratio <= 49:
        return "D"
    else:
        return "E"

def analyze_one_image(tree_id: str, image_path: str):
    if not os.path.exists(image_path):
        logging.info(f"[오류] 파일 없음: {image_path}")
        return None

    img_bgr = cv2.imread(image_path)
    if img_bgr is None:
        logging.info(f"[오류] OpenCV로 읽지 못함: {image_path}")
        return None

    img_hsv = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2HSV)

    outer_mask = np.zeros(img_hsv.shape[:2], dtype=np.uint8)
    for key in ["나무둘레", "나무표시", "나무표시2"]:
        lo, up = color_ranges[key]
        tmp = cv2.inRange(img_hsv, lo, up)
        outer_mask = cv2.bitwise_or(outer_mask, tmp)

    kernel = np.ones((3,3), np.uint8)
    outer_mask = cv2.dilate(outer_mask, kernel, iterations=1)
    outer_mask = cv2.erode(outer_mask, kernel, iterations=1)

    contours, _ = cv2.findContours(outer_mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if not contours:
        logging.info(f"[결과] 외곽 컨투어 없음: {image_path}")
        return None

    largest = max(contours, key=cv2.contourArea)
    roi_mask = np.zeros(img_bgr.shape[:2], dtype=np.uint8)
    cv2.drawContours(roi_mask, [largest], -1, 255, -1)

    target_colors = ["검정", "갈색", "초록", "보라", "파랑"]
    color_counts = {}
    for c in target_colors:
        lo, up = color_ranges[c]
        mask_c = cv2.inRange(img_hsv, lo, up)
        final_mask = cv2.bitwise_and(mask_c, mask_c, mask=roi_mask)
        cnt = cv2.countNonZero(final_mask)
        color_counts[c] = cnt

    sum_of_5 = sum(color_counts.values())
    if sum_of_5 == 0:
        logging.info(f"[결과] 5색 픽셀 없음: {image_path}")
        return None

    black_brown = color_counts["검정"] + color_counts["갈색"]
    black_brown_ratio = round((black_brown / sum_of_5)*100, 2)

    gpb = color_counts["초록"] + color_counts["보라"] + color_counts["파랑"]
    gpb_ratio = round((gpb / sum_of_5)*100, 2)

    overall_grade = calc_grade(gpb_ratio)

    return {
        "tree_id": tree_id,
        "image_path": image_path,
        "color_counts": color_counts,
        "sum_of_5": sum_of_5,
        "black_brown_count": black_brown,
        "black_brown_ratio": black_brown_ratio,
        "green_purple_blue_count": gpb,
        "green_purple_blue_ratio": gpb_ratio,
        "overall_grade": overall_grade
    }

def analyze_multiple_images(image_list, excel_filename="analysis_multi.xlsx"):
    global grade_count
    grade_count = {"A":0, "B":0, "C":0, "D":0, "E":0}

    if os.path.exists(excel_filename):
        os.remove(excel_filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "분석결과"

    ws.merge_cells("A1:G1")
    ws["A1"] = "음파단층촬영 조사현황"
    ws.append(["수목번호", "이미지", "구분", "픽셀수", "합계", "비율(%)", "등급"])

    current_row = 3

    for (tree_id, img_path) in image_list:
        result = analyze_one_image(tree_id, img_path)
        if result is None:
            continue

        color_counts = result["color_counts"]
        sum_of_5     = result["sum_of_5"]
        black_brown_count = result["black_brown_count"]
        black_brown_ratio = result["black_brown_ratio"]
        gpb_count = result["green_purple_blue_count"]
        gpb_ratio = result["green_purple_blue_ratio"]
        overall_grade = result["overall_grade"]

        grade_count[overall_grade] += 1

        start_row = current_row
        end_row   = start_row + 4

        ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)  # A
        ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)  # B
        ws.merge_cells(start_row=start_row, start_column=7, end_row=end_row, end_column=7)  # G

        ws.cell(row=start_row, column=1, value=tree_id)
        ws.cell(row=start_row, column=2, value=img_path)
        ws.cell(row=start_row, column=7, value=overall_grade)

        ws.merge_cells(start_row=start_row, start_column=5, end_row=start_row+1, end_column=5)
        ws.merge_cells(start_row=start_row, start_column=6, end_row=start_row+1, end_column=6)
        ws.cell(row=start_row,   column=5, value=black_brown_count)
        ws.cell(row=start_row,   column=6, value=black_brown_ratio)

        ws.merge_cells(start_row=start_row+2, start_column=5, end_row=end_row, end_column=5)
        ws.merge_cells(start_row=start_row+2, start_column=6, end_row=end_row, end_column=6)
        ws.cell(row=start_row+2, column=5, value=gpb_count)
        ws.cell(row=start_row+2, column=6, value=gpb_ratio)

        c_list = ["검정", "갈색", "초록", "보라", "파랑"]
        for i, cname in enumerate(c_list):
            r = start_row + i
            ws.cell(row=r, column=3, value=cname)
            ws.cell(row=r, column=4, value=color_counts[cname])

        if os.path.exists(img_path):
            try:
                excel_img = XLImage(img_path)
                excel_img.width = 140
                excel_img.height = 140
                anchor_cell = f"B{start_row}"
                ws.add_image(excel_img, anchor_cell)

                row_height_pts = 21
                for rr in range(start_row, end_row+1):
                    ws.row_dimensions[rr].height = row_height_pts

                ws.column_dimensions["B"].width = 17.5

            except Exception as e:
                logging.info(f"[이미지삽입오류] {img_path}: {e}")

        current_row += 5

    ws.cell(row=1, column=10, value="등급")
    ws.cell(row=1, column=11, value="기준")
    ws.cell(row=1, column=12, value="수량")

    grade_table = [
        ("A",   "0",      grade_count["A"]),
        ("B",   "1-19",   grade_count["B"]),
        ("C",   "20-39",  grade_count["C"]),
        ("D",   "40-49",  grade_count["D"]),
        ("E",   "50이상", grade_count["E"]),
    ]
    row2 = 2
    for i, (g, criteria, qty) in enumerate(grade_table):
        rr = row2 + i
        ws.cell(row=rr, column=10, value=g)
        ws.cell(row=rr, column=11, value=criteria)
        ws.cell(row=rr, column=12, value=qty)

    thin_side = Side(style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    last_row_main = current_row - 1
    for row_cells in ws.iter_rows(min_row=1, max_row=last_row_main, min_col=1, max_col=7):
        for cell in row_cells:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

    for row_cells in ws.iter_rows(min_row=1, max_row=6, min_col=10, max_col=12):
        for cell in row_cells:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

    wb.save(excel_filename)
    logging.info(f"[결과] 전체 {len(image_list)}개 이미지 분석 완료 → {excel_filename}")

def decode_and_run(json_str):
    data = json.loads(json_str)
    image_list = []
    for i in range(1, 16):
        num_key  = f"img{i}Num"
        img_key  = f"img{i}"
        if num_key in data and img_key in data:
            tree_id = data[num_key]
            b64_str = data[img_key]
            if not b64_str:
                continue
            local_path = f"temp_img{i}.jpg"
            with open(local_path, "wb") as f:
                f.write(base64.b64decode(b64_str))
            image_list.append((tree_id, local_path))

    if not image_list:
        logging.info("[결과] 디코딩된 이미지가 하나도 없습니다.")
        return
    excel_out = "analysis.xlsx"
    analyze_multiple_images(image_list, excel_out)

# --------------------------------
# (2) Azure Function main() 함수
# --------------------------------
def main(req: func.HttpRequest) -> func.HttpResponse:
    """
    Azure Function의 HTTP Trigger 엔트리 포인트
    """
    logging.info("Python HTTP trigger function processed a request.")

    # 1) HTTP Body 읽기
    try:
        body_str = req.get_body().decode('utf-8')
    except:
        return func.HttpResponse("Invalid request body", status_code=400)

    # 2) decode_and_run 호출
    decode_and_run(body_str)

    # 3) 엑셀 파일을 Base64로 응답 (또는 그냥 "Success" 반환)
    excel_file = "analysis.xlsx"
    if os.path.exists(excel_file):
        with open(excel_file, "rb") as f:
            excel_bytes = f.read()
        # base64로 인코딩
        excel_b64 = base64.b64encode(excel_bytes).decode("utf-8")
        # JSON 형태로 응답
        resp_data = {
            "result": "success",
            "excelBase64": excel_b64
        }
        return func.HttpResponse(
            json.dumps(resp_data),
            status_code=200,
            mimetype="application/json"
        )
    else:
        return func.HttpResponse("No excel output generated.", status_code=200)
